"""Generate efficiency comparison xlsx + LaTeX table snippet."""
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Efficiency Comparison"

header_font = Font(bold=True, size=11)
section_font = Font(bold=True, size=12, color="C55A11")
header_fill = PatternFill(start_color="D9E2F3", end_color="D9E2F3", fill_type="solid")
section_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
thin_border = Border(
    left=Side(style='thin'), right=Side(style='thin'),
    top=Side(style='thin'), bottom=Side(style='thin')
)
center = Alignment(horizontal='center', vertical='center', wrap_text=True)

def write_row(ws, row, values, bold=False, fill=None):
    for c, v in enumerate(values, 1):
        cell = ws.cell(row=row, column=c, value=v)
        cell.alignment = center
        cell.border = thin_border
        if bold:
            cell.font = header_font
        if fill:
            cell.fill = fill

ws.merge_cells('A1:J1')
ws.cell(row=1, column=1, value="Vulnerability Localization & Explanation — Efficiency Comparison (seconds per CVE)").font = Font(bold=True, size=14)
ws.cell(row=1, column=1).alignment = Alignment(horizontal='center')

row = 3
headers = ["Category", "Method", "Data Pre-processing (s/CVE)", "Localization (s/CVE)",
           "Explanation (s/CVE)", "Total (s/CVE)", "GPU Required", "Coverage", "Notes", ""]
write_row(ws, row, headers, bold=True, fill=header_fill)

# Data: [Category, Method, Preprocessing, Localization, Explanation, Total, GPU, Coverage, Notes, ""]
data_rows = [
    ["Our Method", "ExpLoc (Ours)", "~1357", "42.5", "1.6", "~1401", "No", "All GT", "RMS(77s) + CodeQL(779s) + ScubaTrace(~500s) + merged ICFG trace(1.6s)"],

    ["Localization\nBaselines", "RepoMem", "32*", "32.9", "0", "~65", "No", "File-level", "*Single-worker est.; 4.0s w/ 8 workers"],
    ["", "HitMore", "~2†", "~6.6", "0", "~9", "Yes", "File-level", "†MLP train (~120s/886) + feature extraction"],
    ["", "RLocator", "~7‡", "182.3", "0", "~189", "No", "File-level", "‡CodeBERT load (5s/886) + per-CVE file scan"],
    ["", "LocAgent", "~6", "~180", "0", "~186", "No", "File-level", "ICFG load + BM25; multi-turn LLM agent"],

    ["Explanation\nBaselines", "IVDetect + GNNExplainer", "1.6*", "0", "~1.4", "~3.0", "Yes", "Vul-pred (80%)", "*IVDetect train (1424s/886) amortized"],
    ["", "GNN + CFExplainer", "1.5*", "0", "~1.0", "~2.5", "Yes", "Vul-pred (81%)", "*GNN train (1304s/886) amortized"],
    ["", "Coca + GNNExplainer", "2.5*", "0", "~0.7", "~3.2", "Yes", "Vul-pred (66%)", "*Coca train (2206s/886) amortized"],
]

for i, d in enumerate(data_rows):
    fill = section_fill if d[0] else None
    write_row(ws, row + 1 + i, d, fill=fill)

# Column widths
for col, w in {'A': 18, 'B': 22, 'C': 24, 'D': 22, 'E': 22, 'F': 16, 'G': 14, 'H': 18, 'I': 56, 'J': 5}.items():
    ws.column_dimensions[col].width = w

output = "/mnt/d/Desktop/Explainable_Baseline/unified_training/checkpoints/efficiency_v2.xlsx"
wb.save(output)
print(f"Saved to {output}")

# ── Print LaTeX table ──
latex = r"""\begin{table}[htbp]
  \centering
  \caption{Efficiency comparison of vulnerability localization and explanation methods (average seconds per CVE).}
  \resizebox{0.47\textwidth}{!}{
    \begin{tabular}{lccc}
    \toprule
    \multicolumn{1}{c}{Approach} & Data Pre-processing & Vulnerability Localization & Explanation Generation \\
    \midrule
    \tool & 1357     & 42.5 & 1.6 \\
    \midrule
    RepoMem & 32  & 32.9  & -- \\
    HitMore & $\sim$2       & 6.6   & -- \\
    RLocator & $\sim$7       & 182.3 & -- \\
    LocAgent & $\sim$6       & $\sim$180 & -- \\
    \midrule
    IVDetect\textsubscript{exp} & 1.6       & --     & 1.4 \\
    CFExplainer & 1.5       & --     & 1.0 \\
    Coca\textsubscript{exp} & 2.5       & --     & 0.7 \\
    \bottomrule
    \end{tabular}}
  \label{tab:efficiency}%
\end{table}"""

print("\n" + "="*60)
print("LATEX TABLE:")
print(latex)
